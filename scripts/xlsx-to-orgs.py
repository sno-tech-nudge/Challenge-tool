"""One-off: reads the Ecociate org data sheet and merges its organisations into data/orgs.json.
Organisations already in data/orgs.json (by key) are left untouched, so hand-edited content and
Ecociate scores are never overwritten. Usage: python scripts/xlsx-to-orgs.py path/to/sheet.xlsx"""
import json, re, sys
import openpyxl

SKIP_KEYS = {'win-win-bites', 'able-edutech-youvati'}  # already loaded with fuller content
src = sys.argv[1]
ws = openpyxl.load_workbook(src, data_only=True)['Org data']
COL = {ws.cell(5, c).column_letter: c for c in range(1, 25)}


def cell(row, letter):
    v = ws[f'{letter}{row}'].value
    return str(v).strip() if v is not None else ''


def slug(name):
    return re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-')


def lines(text):
    return [re.sub(r'^•\s*', '', l).strip() for l in text.split('\n') if l.strip()]


def not_in_report(text):
    return text.lower().startswith('not in report') or text.lower().startswith('none reported')


def bullets_or_p(text):
    out = []
    for l in lines(text):
        out.append(l)
    return out


def brief(row):
    parts = [p.strip() for p in cell(row, 'B').split('|')]
    entity, typ, started = (parts + ['', '', ''])[:3]
    desc = f"{entity} is registered as a {typ.lower()}. {started}.".replace('..', '.')
    barriers = [b.strip().capitalize() for b in cell(row, 'D').split(';') if b.strip()]
    return [{'kind': 'brief', 'description': desc, 'model': cell(row, 'C'), 'barriers': barriers}]


def geography(row):
    out = []
    states, dists = cell(row, 'E'), cell(row, 'F')
    if states.lower().startswith('reported as'):
        out.append({'kind': 'p', 'text': f'States: {states}.'})
    elif states:
        out.append({'kind': 'names', 'label': 'States', 'items': [s.strip() for s in states.split(',') if s.strip()]})
    if dists.lower().startswith('names not in report'):
        n = re.search(r'\((\d+) districts\)', dists)
        out.append({'kind': 'p', 'text': f'Operational districts: {n.group(1)} (names not in report).' if n else 'Operational districts: names not in report.'})
    elif dists.lower().startswith('not in report'):
        out.append({'kind': 'p', 'text': 'Operational districts: not in report.'})
    elif dists:
        out.append({'kind': 'names', 'label': 'Operational districts', 'items': [d.strip() for d in dists.split(',') if d.strip()]})
    return out


def wfps(row):
    ind = cell(row, 'G')
    col = cell(row, 'H')
    m = re.match(r'(\d+) collectives? with (\d+)', col)
    if m:
        coll = {'title': 'Collectives', 'total': m.group(2), 'lines': [['From collectives', m.group(1)]]}
    else:
        coll = {'title': 'Collectives', 'total': '0', 'lines': []}
    blocks = [{'kind': 'panels', 'panels': [{'title': 'Individuals', 'total': ind, 'lines': []}, coll]}]
    total = cell(row, 'I')
    if '⚠' in total:
        head, note = total.split('⚠', 1)
        blocks.append({'kind': 'p', 'text': f'Total as reported: {head.strip()}. {note.strip()}'})
    return blocks


def markets(row):
    snap = cell(row, 'J').split('\n')
    before = next((s[len('Before:'):].strip() for s in snap if s.startswith('Before:')), '')
    now = next((s[len('Now:'):].strip() for s in snap if s.startswith('Now:')), '')
    out = [{'kind': 'compare', 'before': before, 'now': now}]
    out.append({'kind': 'subhead', 'text': 'Channels the foodpreneurs were already using'})
    k = cell(row, 'K')
    if not_in_report(k):
        out.append({'kind': 'p', 'text': 'None reported.'})
    else:
        out.append({'kind': 'items', 'rows': [{'title': t, 'lines': []} for t in lines(k)]})
    out.append({'kind': 'subhead', 'text': 'Channels opened through the organisation (since intervention)'})
    out.append({'kind': 'items', 'rows': [{'title': t, 'lines': []} for t in lines(cell(row, 'L'))]})
    rev = cell(row, 'O')
    note = None
    if '⚠' in rev:
        rev, note = rev.split('⚠', 1)
    rows = []
    for l in lines(rev):
        label, _, val = l.partition(':')
        rows.append([{'Total': 'Cumulative revenue', 'From new channels': 'From new markets', 'Average per WFP': 'Average per foodpreneur (based on data submitted)'}.get(label.strip(), label.strip()), val.strip()])
    out.append({'kind': 'subhead', 'text': 'Revenue earned by the foodpreneurs (Jan to Jul 2026)'})
    out.append({'kind': 'table', 'head': [k for k, _ in rows], 'rows': [[v for _, v in rows]]})
    if note:
        out.append({'kind': 'p', 'text': note.strip()})
    return out


def partnerships(row):
    out = [{'kind': 'subhead', 'text': 'Active partnerships'}, {'kind': 'bullets', 'items': lines(cell(row, 'P'))}]
    q = cell(row, 'Q')
    if q.lower().strip() != 'not in report.':
        out.append({'kind': 'subhead', 'text': 'Proposed partnerships'})
        if not_in_report(q):
            out.append({'kind': 'p', 'text': q})
        else:
            out.append({'kind': 'bullets', 'items': lines(q)})
    return out


def economics(row):
    return [
        {'kind': 'subhead', 'text': 'Solution model: revenue'},
        {'kind': 'bullets', 'items': lines(cell(row, 'R'))},
    ]


def plan(row):
    return [{'kind': 'p', 'text': 'Targets for August 2026 to May 2027.'}, {'kind': 'bullets', 'items': lines(cell(row, 'T'))}]


def swot(row):
    return [{'kind': 'swot', 's': lines(cell(row, 'U')), 'w': lines(cell(row, 'V')), 'o': lines(cell(row, 'W')), 't': lines(cell(row, 'X'))}]


path = 'data/orgs.json'
orgs = json.load(open(path, encoding='utf-8'))
have = {o['key'] for o in orgs}
added = []
for row in range(6, ws.max_row + 1):
    name = cell(row, 'A')
    if not name:
        continue
    key = slug(name)
    if key in SKIP_KEYS or key in have:
        continue
    orgs.append({
        'key': key,
        'name': name,
        'details': {
            'brief': brief(row), 'geography': geography(row), 'wfps': wfps(row), 'markets': markets(row),
            'partnerships': partnerships(row), 'economics': economics(row), 'plan': plan(row), 'swot': swot(row),
        },
    })
    added.append(name)
json.dump(orgs, open(path, 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
print('added', len(added), added)
